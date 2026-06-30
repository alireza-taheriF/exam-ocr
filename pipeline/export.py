"""Export processed exam blocks to a Persian RTL Word document."""

from __future__ import annotations

import os
import tempfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image

Block = dict[str, Any]

# Persian-friendly font; Tahoma is widely available on most systems.
_BODY_FONT = "Tahoma"
_BODY_FONT_SIZE = Pt(12)
_MAX_IMAGE_WIDTH = Inches(6)
_IMAGE_DPI = 96
_BLOCK_SPACE_BEFORE = Pt(4)
_BLOCK_SPACE_AFTER = Pt(10)
_LINE_SPACING = 1.15


def _set_paragraph_rtl(paragraph) -> None:
    """Mark a paragraph as right-to-left."""
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:bidi")) is None:
        p_pr.append(OxmlElement("w:bidi"))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _set_run_font(run, font_name: str = _BODY_FONT) -> None:
    """Apply body font to a run, including complex-script (Persian) mapping."""
    run.font.name = font_name
    run.font.size = _BODY_FONT_SIZE
    r_pr = run._element.get_or_add_rPr()
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:cs"), font_name)
    r_fonts.set(qn("w:eastAsia"), font_name)
    existing = r_pr.find(qn("w:rFonts"))
    if existing is not None:
        r_pr.remove(existing)
    r_pr.insert(0, r_fonts)


def _apply_block_spacing(paragraph, *, first_block: bool) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0) if first_block else _BLOCK_SPACE_BEFORE
    fmt.space_after = _BLOCK_SPACE_AFTER
    fmt.line_spacing = _LINE_SPACING


def _configure_document_rtl(doc: Document) -> None:
    """Set document defaults for Persian RTL body text."""
    section = doc.sections[0]
    sect_pr = section._sectPr
    if sect_pr.find(qn("w:bidi")) is None:
        sect_pr.append(OxmlElement("w:bidi"))

    normal = doc.styles["Normal"]
    normal.font.name = _BODY_FONT
    normal.font.size = _BODY_FONT_SIZE
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    normal_p_pr = normal._element.get_or_add_pPr()
    if normal_p_pr.find(qn("w:bidi")) is None:
        normal_p_pr.append(OxmlElement("w:bidi"))


def _has_valid_text(block: Block) -> bool:
    if block.get("illegible"):
        return False
    text = block.get("text")
    return bool(text and str(text).strip())


def _picture_display_width(image: Image.Image) -> Inches:
    """Fit image within max page width without upscaling small crops too much."""
    width_inches = image.width / _IMAGE_DPI
    if width_inches > _MAX_IMAGE_WIDTH.inches:
        return _MAX_IMAGE_WIDTH
    return Inches(max(width_inches, 0.5))


def _add_text_block(doc: Document, text: str, *, first_block: bool) -> None:
    paragraph = doc.add_paragraph(str(text).strip())
    _set_paragraph_rtl(paragraph)
    _apply_block_spacing(paragraph, first_block=first_block)
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run(str(text).strip())
    _set_run_font(run, _BODY_FONT)


def _add_crop_block(doc: Document, crop_image: Image.Image, *, first_block: bool) -> None:
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
        crop_image.save(tmp.name)
        tmp_path = tmp.name

    try:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _apply_block_spacing(paragraph, first_block=first_block)
        run = paragraph.add_run()
        run.add_picture(tmp_path, width=_picture_display_width(crop_image))
    except Exception:
        fallback = doc.add_paragraph("[تصویر — قابل نمایش نیست]")
        _set_paragraph_rtl(fallback)
        _apply_block_spacing(fallback, first_block=first_block)
        _set_run_font(fallback.runs[0] if fallback.runs else fallback.add_run(), _BODY_FONT)
    finally:
        os.unlink(tmp_path)


def export_to_word(blocks: list[Block], output_path: str) -> str:
    """
    Interleave typed Persian text and cropped images in layout reading order.

    - Valid OCR text (not illegible) → RTL paragraph
    - crop_image (table, figure, or illegible text) → centered inline image
    """
    doc = Document()
    _configure_document_rtl(doc)

    first_block = True
    wrote_any = False
    for block in blocks:
        wrote_content = False

        if _has_valid_text(block):
            _add_text_block(doc, str(block["text"]), first_block=first_block)
            wrote_content = True

        crop_image = block.get("crop_image")
        if crop_image is not None:
            _add_crop_block(doc, crop_image, first_block=first_block and not wrote_content)
            wrote_content = True

        if wrote_content:
            first_block = False
            wrote_any = True

    if not wrote_any:
        para = doc.add_paragraph("[هیچ محتوای قابل استخراجی یافت نشد]")
        _set_paragraph_rtl(para)
        _set_run_font(para.runs[0] if para.runs else para.add_run(), _BODY_FONT)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return str(output)
